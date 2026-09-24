import os
import re
import sys
import openpyxl

def escape_sql_string(val):
    if val is None:
        return "NULL"
    val_str = str(val).strip()
    if not val_str:
        return "NULL"
    escaped = val_str.replace("'", "''")
    return f"N'{escaped}'"

def format_sql_number(val, default=0):
    if val is None:
        return str(default)
    if isinstance(val, (int, float)):
        return str(int(val))
    val_str = str(val).strip().lower()
    if val_str == "free":
        return "0"
    cleaned = re.sub(r"[^\d]", "", val_str)
    if cleaned:
        return str(int(cleaned))
    return str(default)

def main():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    excel_path = os.path.join(base_dir, "Online Dress Rental Store Plan.xlsx")
    dress_img_dir = os.path.join(base_dir, "DressImages")
    acc_img_dir = os.path.join(base_dir, "AccessoryImages")
    out_sql_path = os.path.join(base_dir, "Insert_Seed_Data.sql")

    print(f"Loading Excel: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    # 1. Categories Mapping
    types = [
        (1, "Váy & Trang phục", "Các mẫu váy dạ hội, váy tiệc, áo dài và đồ đi biển"),
        (2, "Phụ kiện thời trang", "Trang sức, túi xách, giày, khăn choàng và phụ kiện đi kèm")
    ]

    categories = [
        # Type 1
        (1, 1, "Váy dạ hội & Sang trọng", "Váy dạ hội lộng lẫy, sang trọng (Gorgeous & Elegant)"),
        (2, 1, "Váy tiệc & Thiết kế", "Váy dự tiệc, sự kiện, sinh nhật, prom"),
        (3, 1, "Áo khoác lông", "Áo lông, áo choàng ấm áp đi tiệc"),
        (4, 1, "Áo dài & Cổ phục", "Áo dài truyền thống, cổ phục Việt, áo yếm, sườn xám"),
        (5, 1, "Đi biển & Dạo phố", "Trang phục đi biển, dạo phố, chụp ảnh ngoài trời (Picnic & Beach)"),
        # Type 2
        (6, 2, "Dây chuyền", "Vòng cổ, dây chuyền thời trang đính đá, ngọc trai"),
        (7, 2, "Khuyên tai", "Hoa tai, khuyên kẹp, khuyên bạc, khuyên đính đá"),
        (8, 2, "Nhẫn", "Nhẫn thời trang, nhẫn bạc, nhẫn đá"),
        (9, 2, "Vòng tay", "Lắc tay, vòng tay vàng, titan"),
        (10, 2, "Phụ kiện tóc", "Kẹp tóc, cài tóc cô dâu, bờm tóc"),
        (11, 2, "Túi xách", "Túi xách dạ tiệc, ví cầm tay"),
        (12, 2, "Giày", "Giày cao gót, giày quai dây, giày đính nơ"),
        (13, 2, "Găng tay", "Găng tay dạ tiệc, găng tay ren"),
        (14, 2, "Khăn choàng", "Khăn choàng lông, khăn lông cừu"),
        (15, 2, "Kính mắt", "Kính râm, kính thời trang"),
        (16, 2, "Bộ dây chuyền & Nhẫn", "Bộ trang sức kết hợp dây chuyền và nhẫn"),
        (17, 2, "Bộ trang sức", "Set trang sức hoàn chỉnh (dây chuyền, nhẫn, hoa tai)")
    ]

    acc_cat_map = {
        "necklace": 6,
        "earring": 7,
        "ring": 8,
        "bracelet": 9,
        "hair accessory": 10,
        "bag": 11,
        "shoes": 12,
        "gloves": 13,
        "scarf": 14,
        "glasses": 15,
        "necklace/ ring": 16,
        "necklace/ ring/ earring": 17
    }

    # 2. Parse Dress List
    ws_dress = wb["Dress List"]
    dress_products = []
    dress_items = []

    for r in range(2, ws_dress.max_row + 1):
        c_code = ws_dress.cell(r, 6).value
        c_name = ws_dress.cell(r, 5).value
        if c_code is None and c_name is None:
            continue
        code = int(c_code)
        name = str(c_name).strip() if c_name else f"Váy thiết kế {code}"

        # Category determination
        if 1001 <= code <= 1011:
            cat_id = 1
        elif 1012 <= code <= 1086:
            cat_id = 2
        elif 1087 <= code <= 1090:
            cat_id = 3
        elif 1091 <= code <= 1098:
            cat_id = 4
        elif 1099 <= code <= 1113:
            cat_id = 5
        else:
            cat_id = 2

        color = ws_dress.cell(r, 8).value
        orig_p = format_sql_number(ws_dress.cell(r, 9).value, 0)
        p3 = format_sql_number(ws_dress.cell(r, 10).value, 0)
        p7 = format_sql_number(ws_dress.cell(r, 11).value, 0)
        notes = ws_dress.cell(r, 14).value
        size = ws_dress.cell(r, 7).value
        size_str = str(size).strip() if size else "Free"
        status_val = ws_dress.cell(r, 12).value
        
        status = "AVAILABLE"
        if status_val:
            s_lower = str(status_val).strip().lower()
            if "rented" in s_lower:
                status = "RENTED"
            elif "custom" in s_lower:
                status = "IN_REPAIR"

        dress_products.append({
            "id": code,
            "category_id": cat_id,
            "name": name,
            "description": notes,
            "color": str(color).strip() if color else None,
            "original_price": orig_p,
            "price_3d": p3,
            "price_7d": p7,
            "deposit_rate": 100
        })

        dress_items.append({
            "id": code,
            "product_id": code,
            "item_code": str(code),
            "size": size_str,
            "condition": "Mới 99%",
            "status": status
        })

    # 3. Parse Accessory List
    ws_acc = wb["Accesorries List "]
    acc_products = []
    acc_items = []

    for r in range(2, ws_acc.max_row + 1):
        c_code = ws_acc.cell(r, 2).value
        c_name = ws_acc.cell(r, 1).value
        if c_code is None and c_name is None:
            continue
        code = int(c_code)
        
        # Name handling
        if c_name:
            name = str(c_name).strip()
        else:
            if code == 2035:
                name = "Khuyên tai phong cách 2035"
            elif code == 2045:
                name = "Vòng tay thời trang 2045"
            else:
                name = f"Phụ kiện {code}"

        # Category handling
        c_cat = ws_acc.cell(r, 3).value
        if c_cat:
            clean_cat = str(c_cat).strip().lower()
            cat_id = acc_cat_map.get(clean_cat, 7)
        else:
            if code == 2035:
                cat_id = 7 # Earring
            elif code in (2044, 2045):
                cat_id = 9 # Bracelet
            else:
                cat_id = 7

        color = ws_acc.cell(r, 5).value
        if color is None and code in (2044, 2045):
            color = "Gold"

        orig_p = format_sql_number(ws_acc.cell(r, 6).value, 0)
        
        # Rental price handling
        raw_price = ws_acc.cell(r, 7).value
        if raw_price is None and code in (2035, 2044, 2045):
            rent_p = "20000"
        else:
            rent_p = format_sql_number(raw_price, 0)

        notes = ws_acc.cell(r, 10).value
        size = ws_acc.cell(r, 4).value
        size_str = str(size).strip() if size else "Freesize"
        status_val = ws_acc.cell(r, 8).value
        
        status = "AVAILABLE"
        if status_val and "rented" in str(status_val).strip().lower():
            status = "RENTED"

        acc_products.append({
            "id": code,
            "category_id": cat_id,
            "name": name,
            "description": notes,
            "color": str(color).strip() if color else None,
            "original_price": orig_p,
            "price_3d": rent_p,
            "price_7d": rent_p,
            "deposit_rate": 100
        })

        acc_items.append({
            "id": code,
            "product_id": code,
            "item_code": str(code),
            "size": size_str,
            "condition": "Mới 99%",
            "status": status
        })

    # 4. Parse Images
    product_images = []
    img_counter = 1

    # Dresses images
    for p in dress_products:
        p_id = p["id"]
        code_str = str(p_id)
        folder = os.path.join(dress_img_dir, code_str)
        files = []
        if os.path.exists(folder):
            files = sorted([f for f in os.listdir(folder) if os.path.isfile(os.path.join(folder, f))])
        
        # Handle twin empty folders
        if not files:
            if p_id == 1089:
                folder_src = os.path.join(dress_img_dir, "1087")
                if os.path.exists(folder_src):
                    files = sorted([f for f in os.listdir(folder_src) if os.path.isfile(os.path.join(folder_src, f))])
                    folder = folder_src
            elif p_id == 1090:
                folder_src = os.path.join(dress_img_dir, "1088")
                if os.path.exists(folder_src):
                    files = sorted([f for f in os.listdir(folder_src) if os.path.isfile(os.path.join(folder_src, f))])
                    folder = folder_src

        for idx, f in enumerate(files):
            # Use relative forward-slash path
            rel_folder = os.path.basename(os.path.dirname(os.path.join(folder, f)))
            rel_path = f"DressImages/{rel_folder}/{f}"
            is_thumb = 1 if idx == 0 else 0
            product_images.append({
                "id": img_counter,
                "product_id": p_id,
                "image_url": rel_path,
                "is_thumbnail": is_thumb,
                "display_order": idx + 1
            })
            img_counter += 1

    # Accessories images
    for p in acc_products:
        p_id = p["id"]
        code_str = str(p_id)
        folder = os.path.join(acc_img_dir, code_str)
        files = []
        if os.path.exists(folder):
            files = sorted([f for f in os.listdir(folder) if os.path.isfile(os.path.join(folder, f))])

        for idx, f in enumerate(files):
            rel_path = f"AccessoryImages/{code_str}/{f}"
            is_thumb = 1 if idx == 0 else 0
            product_images.append({
                "id": img_counter,
                "product_id": p_id,
                "image_url": rel_path,
                "is_thumbnail": is_thumb,
                "display_order": idx + 1
            })
            img_counter += 1

    print(f"Total Type: {len(types)}")
    print(f"Total Category: {len(categories)}")
    print(f"Total Products: {len(dress_products) + len(acc_products)} ({len(dress_products)} dresses, {len(acc_products)} accessories)")
    print(f"Total Items: {len(dress_items) + len(acc_items)}")
    print(f"Total Images: {len(product_images)}")

    # 5. Generate SQL
    lines = []
    lines.append("-- =============================================================================")
    lines.append("-- SCRIPT INSERT DỮ LIỆU SEED TỪ EXCEL & THƯ MỤC ẢNH VÀO DATABASE ENCHANTED_RENTAL")
    lines.append("-- Được tạo tự động bởi generate_insert_sql.py")
    lines.append("-- =============================================================================")
    lines.append("USE [enchanted_rental];")
    lines.append("GO")
    lines.append("SET NOCOUNT ON;")
    lines.append("GO")
    lines.append("")
    lines.append("BEGIN TRANSACTION;")
    lines.append("BEGIN TRY")
    lines.append("")

    # Type
    lines.append("    -- -------------------------------------------------------------------------")
    lines.append("    -- 1. BẢNG [dbo].[Type]")
    lines.append("    -- -------------------------------------------------------------------------")
    lines.append("    SET IDENTITY_INSERT [dbo].[Type] ON;")
    for t_id, t_name, t_desc in types:
        lines.append(f"    INSERT INTO [dbo].[Type] ([Id], [Name], [Description]) VALUES ({t_id}, {escape_sql_string(t_name)}, {escape_sql_string(t_desc)});")
    lines.append("    SET IDENTITY_INSERT [dbo].[Type] OFF;")
    lines.append("    PRINT N'Inserted [Type] successfully.';")
    lines.append("")

    # Category
    lines.append("    -- -------------------------------------------------------------------------")
    lines.append("    -- 2. BẢNG [dbo].[Category]")
    lines.append("    -- -------------------------------------------------------------------------")
    lines.append("    SET IDENTITY_INSERT [dbo].[Category] ON;")
    for c_id, t_id, c_name, c_desc in categories:
        lines.append(f"    INSERT INTO [dbo].[Category] ([Id], [TypeId], [Name], [Description]) VALUES ({c_id}, {t_id}, {escape_sql_string(c_name)}, {escape_sql_string(c_desc)});")
    lines.append("    SET IDENTITY_INSERT [dbo].[Category] OFF;")
    lines.append("    PRINT N'Inserted [Category] successfully.';")
    lines.append("")

    # Product
    lines.append("    -- -------------------------------------------------------------------------")
    lines.append("    -- 3. BẢNG [dbo].[Product]")
    lines.append("    -- -------------------------------------------------------------------------")
    lines.append("    SET IDENTITY_INSERT [dbo].[Product] ON;")
    all_products = dress_products + acc_products
    for p in all_products:
        lines.append(
            f"    INSERT INTO [dbo].[Product] ([Id], [CategoryId], [Name], [Description], [Color], [OriginalPrice], [Price3Days], [Price7Days], [DepositRate]) "
            f"VALUES ({p['id']}, {p['category_id']}, {escape_sql_string(p['name'])}, {escape_sql_string(p['description'])}, {escape_sql_string(p['color'])}, {p['original_price']}, {p['price_3d']}, {p['price_7d']}, {p['deposit_rate']});"
        )
    lines.append("    SET IDENTITY_INSERT [dbo].[Product] OFF;")
    lines.append("    PRINT N'Inserted [Product] successfully.';")
    lines.append("")

    # ProductItem
    lines.append("    -- -------------------------------------------------------------------------")
    lines.append("    -- 4. BẢNG [dbo].[ProductItem]")
    lines.append("    -- -------------------------------------------------------------------------")
    lines.append("    SET IDENTITY_INSERT [dbo].[ProductItem] ON;")
    all_items = dress_items + acc_items
    for itm in all_items:
        lines.append(
            f"    INSERT INTO [dbo].[ProductItem] ([Id], [ProductId], [ItemCode], [Size], [CurrentCondition], [Status]) "
            f"VALUES ({itm['id']}, {itm['product_id']}, '{itm['item_code']}', '{itm['size']}', {escape_sql_string(itm['condition'])}, '{itm['status']}');"
        )
    lines.append("    SET IDENTITY_INSERT [dbo].[ProductItem] OFF;")
    lines.append("    PRINT N'Inserted [ProductItem] successfully.';")
    lines.append("")

    # ProductImage
    lines.append("    -- -------------------------------------------------------------------------")
    lines.append("    -- 5. BẢNG [dbo].[ProductImage]")
    lines.append("    -- -------------------------------------------------------------------------")
    lines.append("    SET IDENTITY_INSERT [dbo].[ProductImage] ON;")
    for img in product_images:
        lines.append(
            f"    INSERT INTO [dbo].[ProductImage] ([Id], [ProductId], [ImageUrl], [IsThumbnail], [DisplayOrder]) "
            f"VALUES ({img['id']}, {img['product_id']}, {escape_sql_string(img['image_url'])}, {img['is_thumbnail']}, {img['display_order']});"
        )
    lines.append("    SET IDENTITY_INSERT [dbo].[ProductImage] OFF;")
    lines.append("    PRINT N'Inserted [ProductImage] successfully.';")
    lines.append("")

    # Reseed
    lines.append("    -- -------------------------------------------------------------------------")
    lines.append("    -- 6. CẬP NHẬT LẠI GIÁ TRỊ IDENTITY (RESEED)")
    lines.append("    -- -------------------------------------------------------------------------")
    max_type_id = max(t[0] for t in types)
    max_cat_id = max(c[0] for c in categories)
    max_prod_id = max(p["id"] for p in all_products)
    max_item_id = max(itm["id"] for itm in all_items)
    max_img_id = max(img["id"] for img in product_images)

    lines.append(f"    DBCC CHECKIDENT ('[dbo].[Type]', RESEED, {max_type_id});")
    lines.append(f"    DBCC CHECKIDENT ('[dbo].[Category]', RESEED, {max_cat_id});")
    lines.append(f"    DBCC CHECKIDENT ('[dbo].[Product]', RESEED, {max_prod_id});")
    lines.append(f"    DBCC CHECKIDENT ('[dbo].[ProductItem]', RESEED, {max_item_id});")
    lines.append(f"    DBCC CHECKIDENT ('[dbo].[ProductImage]', RESEED, {max_img_id});")
    lines.append("")
    lines.append("    COMMIT TRANSACTION;")
    lines.append("    PRINT N'======================================================';")
    lines.append("    PRINT N'TOÀN BỘ DỮ LIỆU ĐÃ ĐƯỢC INSERT THÀNH CÔNG VÀO DATABASE!';")
    lines.append("    PRINT N'======================================================';")
    lines.append("END TRY")
    lines.append("BEGIN CATCH")
    lines.append("    IF @@TRANCOUNT > 0")
    lines.append("        ROLLBACK TRANSACTION;")
    lines.append("    PRINT N'CÓ LỖI XẢY RA TRONG QUÁ TRÌNH INSERT DỮ LIỆU: ' + ERROR_MESSAGE();")
    lines.append("    THROW;")
    lines.append("END CATCH;")
    lines.append("GO")

    with open(out_sql_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))

    print(f"\nSuccessfully generated SQL script at: {out_sql_path}")
    print(f"Total SQL lines: {len(lines)}")

if __name__ == "__main__":
    main()

